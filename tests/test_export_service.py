"""Pruebas del servicio de exportación (app/services/export_service.py)."""
import csv
import io

import pytest

from app.services.export_service import export_csv, export_excel, export_pdf

SAMPLE_SIM = {
    "name": "Finca de prueba",
    "price_source": "actual",
    "coffee_type": "pergamino_seco",
    "hectares": 5,
    "production_per_hectare_kg": 2000,
    "price_per_kg": 11200,
    "costs": {"mano_de_obra": 15_000_000, "fertilizantes": 8_000_000},
    "total_investment": 23_000_000,
    "total_production_kg": 10000,
    "total_income": 112_000_000,
    "profit": 89_000_000,
    "roi_pct": 386.9,
    "margin_pct": 79.5,
    "break_even_kg": 2053.6,
    "scenarios": [
        {"label": "Esperado", "total_investment": 23_000_000, "total_income": 112_000_000, "profit": 89_000_000, "roi_pct": 386.9},
    ],
}


def test_export_csv_contiene_las_secciones_pedidas():
    buf = export_csv(SAMPLE_SIM)
    content = buf.getvalue()
    for seccion in ["Datos introducidos", "Costos", "Resultados", "Advertencia"]:
        assert seccion in content
    # Debe ser CSV válido y parseable
    reader = csv.reader(io.StringIO(content))
    rows = list(reader)
    assert len(rows) > 5


def test_export_excel_genera_workbook_valido():
    from openpyxl import load_workbook

    buf = export_excel(SAMPLE_SIM)
    wb = load_workbook(buf)
    ws = wb.active
    assert ws["A3"].value == "Nombre"
    assert ws["B3"].value == SAMPLE_SIM["name"]


def test_export_pdf_genera_pdf_valido():
    from pypdf import PdfReader

    buf = export_pdf(SAMPLE_SIM)
    reader = PdfReader(buf)
    assert len(reader.pages) >= 1
    text = reader.pages[0].extract_text()
    assert "Coffee Market Intelligence" in text
